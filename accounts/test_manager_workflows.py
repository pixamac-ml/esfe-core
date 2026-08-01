"""Tests critiques pour le dashboard gestionnaire — workflows métier."""

from html import unescape
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import json
import re
from typing import Any, cast
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from branches.models import Branch
from admissions.models import Candidature, CandidatureDocument
from academics.models import AcademicClass, AcademicYear
from inscriptions.models import Inscription
from formations.models import Programme, Cycle, Diploma, Filiere, RequiredDocument
from accounts.models import (
    BranchCashMovement,
    BranchBankTransfer,
    BranchExpense,
    BranchMonthlyClosure,
    Donation,
    PayrollEntry,
    TeacherHonorariumEntry,
    Profile,
)
from payments.models import CashPaymentSession, Payment, PaymentAgent
from accounts.services.manager_intelligence import reconcile_branch_financial_movements


User = get_user_model()
USER_MANAGER = cast(Any, User._default_manager)


def _create_user(username, groups=None, branch=None, position="branch_manager", **kw):
    user = USER_MANAGER.create_user(
        username=username, email=f"{username}@test.com", password="pass1234", **kw,
    )
    if groups:
        for name in groups:
            g, _ = Group.objects.get_or_create(name=name)
            user.groups.add(g)
    if branch:
        profile = user.profile
        profile.branch = branch
        profile.position = position
        profile.save(update_fields=["branch", "position", "updated_at"])
    return user


def _create_branch(code="TST", name="Test Annexe"):
    return Branch.objects.create(name=name, code=code, slug=f"annexe-{code.lower()}")


def _login(client, user):
    client.force_login(user)


def _create_programme():
    cycle = Cycle.objects.create(name="Licence", min_duration_years=3, max_duration_years=4)
    diploma = Diploma.objects.create(name="Licence Pro", level="superieur")
    filiere = Filiere.objects.create(name="Informatique")
    return Programme.objects.create(
        title="LP Informatique",
        filiere=filiere,
        cycle=cycle,
        diploma_awarded=diploma,
        duration_years=3,
        short_description="Formation en informatique",
        description="Formation en informatique de niveau licence",
    )


def _create_academic_class(programme, branch, level="L1"):
    year, _ = AcademicYear.objects.get_or_create(
        name="2025-2026",
        defaults={"start_date": date(2025, 10, 1), "end_date": date(2026, 7, 31), "is_active": True},
    )
    return AcademicClass.objects.create(programme=programme, branch=branch, academic_year=year, level=level)


def _create_candidature(programme, branch, **kw):
    defaults = dict(
        first_name="Jean",
        last_name="Test",
        email="jean@test.com",
        phone="70000000",
        birth_date=date(2000, 1, 1),
        birth_place="Bamako",
        gender="male",
        academic_year="2025-2026",
    )
    defaults.update(kw)
    return Candidature.objects.create(programme=programme, branch=branch, **defaults)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    CHANNEL_LAYERS={"default": {"BACKEND": "channels.layers.InMemoryChannelLayer"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerCandidatureWorkflowTests(TestCase):
    """Workflow candidature : soumettre → analyser → accepter/rejeter."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_cand", groups=["gestionnaire"], branch=self.branch)
        self.programme = _create_programme()
        self.candidature = _create_candidature(self.programme, self.branch, status="submitted")
        _login(self.client, self.manager)

    def test_candidature_under_review(self):
        url = reverse("accounts:htmx_candidature_under_review", args=[self.candidature.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.candidature.refresh_from_db()
        self.assertEqual(self.candidature.status, "under_review")
        self.assertIn("HX-Trigger", response.headers)
        self.assertIn("showToast", response.headers["HX-Trigger"])

    def test_candidature_accept(self):
        self.candidature.status = "under_review"
        self.candidature.save()
        url = reverse("accounts:htmx_candidature_accept", args=[self.candidature.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.candidature.refresh_from_db()
        self.assertEqual(self.candidature.status, "accepted")
        self.assertIn("showToast", response.headers["HX-Trigger"])

    def test_candidature_reject_with_reason(self):
        url = reverse("accounts:htmx_candidature_reject", args=[self.candidature.id])
        response = self.client.post(url, {"reason": "Dossier incomplet"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.candidature.refresh_from_db()
        self.assertEqual(self.candidature.status, "rejected")
        self.assertIn("Dossier incomplet", self.candidature.rejection_reason)
        self.assertIn("showToast", response.headers["HX-Trigger"])

    def test_candidature_delete_only_when_rejected(self):
        """Ne peut supprimer qu'une candidature rejetee."""
        url = reverse("accounts:htmx_candidature_delete", args=[self.candidature.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 400)

    def test_candidature_delete_when_rejected(self):
        self.candidature.status = "rejected"
        self.candidature.rejection_reason = "Test"
        self.candidature.save()
        url = reverse("accounts:htmx_candidature_delete", args=[self.candidature.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerPaymentWorkflowTests(TestCase):
    """Workflow paiement : valider, annuler."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_pay", groups=["gestionnaire"], branch=self.branch)
        self.programme = _create_programme()
        self.candidature = _create_candidature(self.programme, self.branch, status="accepted")
        self.inscription = Inscription.objects.create(
            candidature=self.candidature,
            amount_due=500000,
            amount_paid=0,
        )
        self.inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        self.inscription.save(update_fields=["status"])
        PaymentAgent.objects.create(
            user=self.manager, branch=self.branch, agent_code="AGT001", is_active=True,
        )
        self.payment = Payment.objects.create(
            inscription=self.inscription,
            amount=100000,
            method=Payment.METHOD_CASH,
        )
        _login(self.client, self.manager)

    def test_payment_validate_credits_caisse(self):
        url = reverse("accounts:htmx_payment_validate", args=[self.payment.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.status, "validated")
        self.assertIn("showToast", response.headers["HX-Trigger"])
        movements = BranchCashMovement.objects.filter(
            branch=self.branch, source=BranchCashMovement.SOURCE_STUDENT_PAYMENT,
        )
        self.assertEqual(movements.count(), 1)
        self.assertEqual(movements.first().amount, 100000)

    def test_payment_cancel(self):
        url = reverse("accounts:htmx_payment_cancel", args=[self.payment.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.status, "cancelled")
        self.assertIn("showToast", response.headers["HX-Trigger"])


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerExpenseWorkflowTests(TestCase):
    """Workflow depense : creer, approuver, payer."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_exp", groups=["gestionnaire"], branch=self.branch)
        self.expense = BranchExpense.objects.create(
            branch=self.branch,
            title="Fournitures bureau",
            category="supplies",
            amount=50000,
            created_by=self.manager,
        )
        _login(self.client, self.manager)

    def test_expense_approve_changes_status(self):
        url = reverse("accounts:htmx_manager_expense_approve", args=[self.expense.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertIn("HX-Redirect", response.headers)
        self.expense.refresh_from_db()
        self.assertEqual(self.expense.status, "approved")

    def _seed_cash(self, amount=100000):
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_MANUAL,
            amount=amount,
            label="Solde initial",
            created_by=self.manager,
        )

    def test_expense_pay_creates_cash_movement(self):
        self._seed_cash()
        self.expense.status = BranchExpense.STATUS_APPROVED
        self.expense.save()
        url = reverse("accounts:htmx_manager_expense_pay", args=[self.expense.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.expense.refresh_from_db()
        self.assertEqual(self.expense.status, "paid")
        movements = BranchCashMovement.objects.filter(
            branch=self.branch, source=BranchCashMovement.SOURCE_EXPENSE,
        )
        self.assertEqual(movements.count(), 1)
        self.assertEqual(movements.first().amount, 50000)

    def test_expense_pay_rejected_when_caisse_insufficient(self):
        self.expense.status = BranchExpense.STATUS_APPROVED
        self.expense.save(update_fields=["status", "updated_at"])
        response = self.client.post(
            reverse("accounts:htmx_manager_expense_pay", args=[self.expense.id]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.expense.refresh_from_db()
        self.assertEqual(self.expense.status, BranchExpense.STATUS_APPROVED)
        self.assertFalse(
            BranchCashMovement.objects.filter(
                branch=self.branch,
                source=BranchCashMovement.SOURCE_EXPENSE,
            ).exists()
        )


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerDonationWorkflowTests(TestCase):
    """Workflow don : creer → mouvement caisse + HX-Trigger."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_don", groups=["gestionnaire"], branch=self.branch)
        _login(self.client, self.manager)

    def test_donation_creates_cash_movement_and_toast(self):
        url = reverse("accounts:htmx_manager_donation_create")
        response = self.client.post(url, {
            "donor_name": "Bienfaiteur X",
            "amount": 100000,
            "date": date.today().isoformat(),
            "motif": "mecenat",
            "payment_method": "cash",
        }, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        donation = Donation.objects.get(branch=self.branch, amount=100000)
        self.assertIsNotNone(donation.cash_movement_id)
        self.assertTrue(
            BranchCashMovement.objects.filter(
                pk=donation.cash_movement_id,
                branch=self.branch,
                source=BranchCashMovement.SOURCE_DONATION,
                amount=100000,
            ).exists(),
        )
        self.assertIn("showToast", response.headers["HX-Trigger"])


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerAccessControlTests(TestCase):
    """Teste les permissions et le filtrage branche."""

    def setUp(self):
        self.branch_a = _create_branch("BRA", "Branche A")
        self.branch_b = _create_branch("BRB", "Branche B")
        self.manager_a = _create_user("mgr_a", groups=["gestionnaire"], branch=self.branch_a)
        self.manager_b = _create_user("mgr_b", groups=["gestionnaire"], branch=self.branch_b)
        self.programme = _create_programme()
        self.cand_a = _create_candidature(
            self.programme, self.branch_a,
            first_name="Alice", last_name="A", email="alice@a.com", phone="70000002",
        )
        self.cand_b = _create_candidature(
            self.programme, self.branch_b,
            first_name="Bob", last_name="B", email="bob@b.com", phone="70000003",
        )

    def test_manager_a_cannot_see_branch_b_candidatures(self):
        _login(self.client, self.manager_a)
        response = self.client.get(
            reverse("accounts:manager_dashboard"),
            {"section": "candidatures"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Alice")
        self.assertNotContains(response, "Bob")

    def test_unauthorized_user_gets_403(self):
        stranger = _create_user("stranger")
        _login(self.client, stranger)
        url = reverse("accounts:htmx_candidature_under_review", args=[self.cand_a.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 403)

    def test_manager_a_cannot_inscribe_candidature_into_branch_b_class(self):
        """Un gestionnaire de l'annexe A ne peut pas inscrire un candidat dans une classe de l'annexe B."""
        self.cand_a.status = "accepted"
        self.cand_a.save(update_fields=["status"])
        class_b = _create_academic_class(self.programme, self.branch_b)
        _login(self.client, self.manager_a)
        url = reverse("accounts:htmx_inscription_create", args=[self.cand_a.id])
        response = self.client.post(
            url,
            {"academic_level": "L1", "academic_class": str(class_b.pk)},
            HTTP_HX_REQUEST="true",
        )
        self.assertFalse(
            Inscription.objects.filter(candidature=self.cand_a).exists(),
            "Une inscription ne doit pas être créée avec une classe d'une autre annexe.",
        )
        self.assertNotEqual(response.status_code, 201)

    def test_manager_a_cannot_access_branch_b_candidature_detail(self):
        """Un gestionnaire de l'annexe A obtient 404 sur la candidature de l'annexe B."""
        _login(self.client, self.manager_a)
        url = reverse("accounts:htmx_candidature_detail", args=[self.cand_b.id])
        response = self.client.get(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 404)

    def test_manager_a_cannot_approve_branch_b_expense(self):
        expense = BranchExpense.objects.create(
            branch=self.branch_b,
            title="Depense annexe B",
            amount=10000,
        )
        _login(self.client, self.manager_a)
        response = self.client.post(
            reverse("accounts:htmx_manager_expense_approve", args=[expense.id]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 404)

    def test_manager_a_cannot_download_branch_b_candidature_document(self):
        document_type = RequiredDocument.objects.create(name="Piece securisee")
        document = CandidatureDocument.objects.create(
            candidature=self.cand_b,
            document_type=document_type,
            file="candidatures/documents/piece-securisee.pdf",
        )
        _login(self.client, self.manager_a)
        response = self.client.get(
            reverse("accounts:manager_candidature_document_download", args=[document.id])
        )
        self.assertEqual(response.status_code, 404)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ManagerPostOnlySecurityTests(TestCase):
    def setUp(self):
        self.branch = _create_branch("PST", "Annexe POST")
        self.manager = _create_user("mgr_post", groups=["gestionnaire"], branch=self.branch)
        self.expense = BranchExpense.objects.create(
            branch=self.branch,
            title="Depense test POST",
            amount=10000,
        )
        _login(self.client, self.manager)

    def test_sensitive_manager_actions_reject_get(self):
        urls = [
            reverse("accounts:htmx_manager_expense_approve", args=[self.expense.id]),
            reverse("accounts:htmx_manager_expense_pay", args=[self.expense.id]),
            reverse("accounts:htmx_manager_monthly_closure_create"),
            reverse("accounts:htmx_manager_donation_create"),
            reverse("accounts:htmx_manager_cash_sync"),
            reverse("accounts:htmx_manager_salary_prepare_all"),
            reverse("accounts:htmx_manager_teacher_honorarium_prepare_all"),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 405)

    def test_manual_cash_entry_cannot_impersonate_automatic_source(self):
        response = self.client.post(
            reverse("accounts:htmx_manager_cash_movement_create"),
            {
                "movement_type": BranchCashMovement.TYPE_IN,
                "source": BranchCashMovement.SOURCE_STUDENT_PAYMENT,
                "amount": "10000",
                "label": "Faux paiement manuel",
                "movement_date": date.today().isoformat(),
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(BranchCashMovement.objects.filter(branch=self.branch).exists())


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerDashboardViewTests(TestCase):
    """Teste le rendu du dashboard principal."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_view", groups=["gestionnaire"], branch=self.branch)
        _login(self.client, self.manager)

    def test_overview_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "overview"})
        self.assertEqual(response.status_code, 200)

    def test_candidatures_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "candidatures"})
        self.assertEqual(response.status_code, 200)

    def test_inscriptions_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "inscriptions"})
        self.assertEqual(response.status_code, 200)

    def test_paiements_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "paiements"})
        self.assertEqual(response.status_code, 200)

    def test_depenses_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "depenses"})
        self.assertEqual(response.status_code, 200)

    def test_caisse_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "caisse"})
        self.assertEqual(response.status_code, 200)

    def test_rapport_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "rapport"})
        self.assertEqual(response.status_code, 200)

    def test_cloture_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "cloture"})
        self.assertEqual(response.status_code, 200)

    def test_dons_section_renders(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "dons"})
        self.assertEqual(response.status_code, 200)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerExportReportTests(TestCase):
    """Teste l'export Excel."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_xls", groups=["gestionnaire"], branch=self.branch)
        _login(self.client, self.manager)

    def test_export_xlsx_returns_file(self):
        url = reverse("accounts:manager_export_report_xlsx")
        response = self.client.get(f"{url}?section=rapport&report_period=month")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        values = {
            str(cell.value)
            for row in workbook["Rapport"].iter_rows()
            for cell in row
            if cell.value is not None
        }
        self.assertIn("Versements bancaires", values)
        self.assertIn("Coupons appliques", values)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    CHANNEL_LAYERS={"default": {"BACKEND": "channels.layers.InMemoryChannelLayer"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerSalaryWorkflowTests(TestCase):
    """Workflow salaires : preparer, payer, avance, notifier."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_sal", groups=["gestionnaire"], branch=self.branch)
        profile = self.manager.profile
        profile.salary_base = 300000
        profile.save(update_fields=["salary_base", "updated_at"])
        self.period_month = date.today().replace(day=1)
        _login(self.client, self.manager)

    def _seed_cash(self, amount=500000):
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_MANUAL,
            amount=amount,
            label="Solde initial",
            created_by=self.manager,
        )

    def test_salary_prepare_all_creates_entry(self):
        url = reverse("accounts:htmx_manager_salary_prepare_all")
        response = self.client.post(
            url, {"period_month": self.period_month.isoformat()}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        entry = PayrollEntry.objects.get(
            branch=self.branch, employee=self.manager, period_month=self.period_month,
        )
        self.assertEqual(entry.base_salary, 300000)
        self.assertEqual(entry.status, PayrollEntry.STATUS_DRAFT)

    def test_salary_pay_creates_cash_movement(self):
        self._seed_cash(500000)
        entry = PayrollEntry.objects.get(
            branch=self.branch, employee=self.manager, period_month=self.period_month,
        )
        entry.status = PayrollEntry.STATUS_READY
        entry.save()
        url = reverse("accounts:htmx_manager_salary_pay", args=[entry.id])
        response = self.client.post(url, {"payment_amount": "100000"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        entry.refresh_from_db()
        self.assertEqual(entry.paid_amount, 100000)
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch, source=BranchCashMovement.SOURCE_PAYROLL, amount=100000,
            ).exists()
        )

    def test_salary_pay_rejected_when_caisse_insufficient(self):
        entry = PayrollEntry.objects.get(
            branch=self.branch, employee=self.manager, period_month=self.period_month,
        )
        entry.status = PayrollEntry.STATUS_READY
        entry.save()
        url = reverse("accounts:htmx_manager_salary_pay", args=[entry.id])
        response = self.client.post(url, {"payment_amount": "100000"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 400)
        entry.refresh_from_db()
        self.assertEqual(entry.paid_amount, 0)

    def test_salary_advance_creates_draft_entry_and_cash_movement(self):
        self._seed_cash(500000)
        url = reverse("accounts:htmx_manager_salary_advance", args=[self.manager.id])
        response = self.client.post(
            url,
            {"advance_amount": "50000", "period_month": self.period_month.isoformat()},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        entry = PayrollEntry.objects.get(
            branch=self.branch, employee=self.manager, period_month=self.period_month,
        )
        self.assertEqual(entry.advances, 50000)
        self.assertEqual(entry.status, PayrollEntry.STATUS_DRAFT)
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch, source=BranchCashMovement.SOURCE_PAYROLL, amount=50000,
            ).exists()
        )

    def test_salary_pay_ready_all_notifies_employee(self):
        entry = PayrollEntry.objects.get(
            branch=self.branch, employee=self.manager, period_month=self.period_month,
        )
        entry.status = PayrollEntry.STATUS_READY
        entry.save()
        url = reverse("accounts:htmx_manager_salary_pay_ready_all")
        response = self.client.post(
            url, {"period_month": self.period_month.isoformat()}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("salaires_disponibles_1", response.headers["HX-Redirect"])


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    CHANNEL_LAYERS={"default": {"BACKEND": "channels.layers.InMemoryChannelLayer"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerHonorariumWorkflowTests(TestCase):
    """Workflow honoraires enseignants : preparer, payer."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_hon", groups=["gestionnaire"], branch=self.branch)
        self.teacher = _create_user("teacher_hon", branch=self.branch, position="teacher")
        profile = self.teacher.profile
        profile.teacher_hourly_rate = 5000
        profile.save(update_fields=["teacher_hourly_rate", "updated_at"])
        self.period_month = date.today().replace(day=1)
        _login(self.client, self.manager)

    def _seed_cash(self, amount=500000):
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_MANUAL,
            amount=amount,
            label="Solde initial",
            created_by=self.manager,
        )

    def test_teacher_honorarium_prepare_all_creates_entry(self):
        url = reverse("accounts:htmx_manager_teacher_honorarium_prepare_all")
        response = self.client.post(
            url, {"period_month": self.period_month.isoformat()}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        entry = TeacherHonorariumEntry.objects.get(
            branch=self.branch, teacher=self.teacher, period_month=self.period_month,
        )
        self.assertEqual(entry.hourly_rate, 5000)
        self.assertEqual(entry.status, TeacherHonorariumEntry.STATUS_DRAFT)

    def test_teacher_honorarium_pay_creates_cash_movement(self):
        self._seed_cash(500000)
        entry = TeacherHonorariumEntry.objects.get(
            branch=self.branch, teacher=self.teacher, period_month=self.period_month,
        )
        entry.validated_hours = Decimal("40")
        entry.status = TeacherHonorariumEntry.STATUS_READY
        entry.save()
        url = reverse("accounts:htmx_manager_teacher_honorarium_pay", args=[entry.id])
        response = self.client.post(url, {"payment_amount": "100000"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        entry.refresh_from_db()
        self.assertEqual(entry.paid_amount, 100000)
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch, source=BranchCashMovement.SOURCE_HONORARIUM, amount=100000,
            ).exists()
        )

    def test_teacher_honorarium_pay_rejected_when_caisse_insufficient(self):
        entry = TeacherHonorariumEntry.objects.get(
            branch=self.branch, teacher=self.teacher, period_month=self.period_month,
        )
        entry.validated_hours = Decimal("40")
        entry.status = TeacherHonorariumEntry.STATUS_READY
        entry.save()
        response = self.client.post(
            reverse("accounts:htmx_manager_teacher_honorarium_pay", args=[entry.id]),
            {"payment_amount": "100000"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        entry.refresh_from_db()
        self.assertEqual(entry.paid_amount, 0)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerMonthlyClosureWorkflowTests(TestCase):
    """Workflow cloture mensuelle : brouillon -> validee -> cloturee."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_clo", groups=["gestionnaire"], branch=self.branch)
        self.period_month = date.today().replace(day=1)
        _login(self.client, self.manager)

    def _create_closure(self):
        url = reverse("accounts:htmx_manager_monthly_closure_create")
        response = self.client.post(
            url,
            {"period_month": self.period_month.isoformat(), "notes": ""},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        return BranchMonthlyClosure.objects.get(branch=self.branch, period_month=self.period_month)

    def test_closure_create_is_draft(self):
        closure = self._create_closure()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_DRAFT)

    def test_closure_validate_from_draft(self):
        closure = self._create_closure()
        url = reverse("accounts:htmx_manager_monthly_closure_validate", args=[closure.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        closure.refresh_from_db()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_VALIDATED)
        self.assertEqual(closure.validated_by, self.manager)
        self.assertIsNotNone(closure.validated_at)

    def test_closure_close_from_validated(self):
        closure = self._create_closure()
        closure.status = BranchMonthlyClosure.STATUS_VALIDATED
        closure.validated_by = self.manager
        closure.validated_at = timezone.now()
        closure.save(update_fields=["status", "validated_by", "validated_at", "updated_at"])
        url = reverse("accounts:htmx_manager_monthly_closure_close", args=[closure.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        closure.refresh_from_db()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_CLOSED)
        self.assertIsNotNone(closure.closed_at)

    def test_closure_validate_fails_if_not_draft(self):
        closure = self._create_closure()
        closure.status = BranchMonthlyClosure.STATUS_CLOSED
        closure.save(update_fields=["status", "updated_at"])
        url = reverse("accounts:htmx_manager_monthly_closure_validate", args=[closure.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertIn("cloture_non_brouillon", response.headers["HX-Redirect"])
        closure.refresh_from_db()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_CLOSED)

    def test_closure_close_fails_if_not_validated(self):
        closure = self._create_closure()
        url = reverse("accounts:htmx_manager_monthly_closure_close", args=[closure.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertIn("cloture_non_validee", response.headers["HX-Redirect"])
        closure.refresh_from_db()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_DRAFT)

    def test_closure_validate_blocks_orphan_donation(self):
        closure = self._create_closure()
        Donation.objects.create(
            branch=self.branch,
            donor_name="Don orphelin",
            amount=25000,
            date=self.period_month,
            created_by=self.manager,
        )
        url = reverse("accounts:htmx_manager_monthly_closure_validate", args=[closure.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Cloture bloquee", response.content.decode())
        closure.refresh_from_db()
        self.assertEqual(closure.status, BranchMonthlyClosure.STATUS_DRAFT)

    def test_closure_create_blocks_orphan_operation(self):
        Donation.objects.create(
            branch=self.branch,
            donor_name="Don orphelin avant cloture",
            amount=15000,
            date=self.period_month,
            created_by=self.manager,
        )
        response = self.client.post(
            reverse("accounts:htmx_manager_monthly_closure_create"),
            {"period_month": self.period_month.isoformat(), "notes": ""},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(
            BranchMonthlyClosure.objects.filter(
                branch=self.branch, period_month=self.period_month,
            ).exists()
        )

    def test_bank_transfer_above_available_cash_is_rejected(self):
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_MANUAL,
            amount=10000,
            label="Petite caisse",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        response = self.client.post(
            reverse("accounts:htmx_manager_monthly_closure_create"),
            {
                "period_month": self.period_month.isoformat(),
                "bank_transfer_amount": "20000",
                "bank_name": "Banque Test",
                "reference": "VIR-EXCESSIF",
                "transfer_date": self.period_month.isoformat(),
                "amount": "20000",
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "depasse la caisse disponible", status_code=400)
        self.assertFalse(BranchBankTransfer.objects.filter(branch=self.branch).exists())

    def test_closure_validate_rejects_negative_cash(self):
        closure = self._create_closure()
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_OUT,
            source=BranchCashMovement.SOURCE_ADJUSTMENT,
            amount=10000,
            label="Anomalie historique",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        response = self.client.post(
            reverse("accounts:htmx_manager_monthly_closure_validate", args=[closure.id]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "caisse de l'annexe est negative", status_code=400)

    def test_closure_recreate_fails_if_not_draft(self):
        closure = self._create_closure()
        closure.status = BranchMonthlyClosure.STATUS_VALIDATED
        closure.save(update_fields=["status", "updated_at"])
        url = reverse("accounts:htmx_manager_monthly_closure_create")
        response = self.client.post(
            url,
            {"period_month": self.period_month.isoformat(), "notes": ""},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerFinancialReconciliationTests(TestCase):
    """Reconciliation des orphelins financiers."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_rec", groups=["gestionnaire"], branch=self.branch)
        self.period_month = date.today().replace(day=1)

    def test_reconciliation_repairs_donation_and_bank_transfer_movements(self):
        donation = Donation.objects.create(
            branch=self.branch,
            donor_name="Donateur Reprise",
            amount=75000,
            date=self.period_month,
            created_by=self.manager,
        )
        closure = BranchMonthlyClosure.objects.create(
            branch=self.branch,
            period_month=self.period_month,
            created_by=self.manager,
        )
        transfer = BranchBankTransfer.objects.create(
            branch=self.branch,
            closure=closure,
            bank_name="Bank X",
            reference="BT-2026-0001",
            transfer_date=self.period_month,
            amount=50000,
            created_by=self.manager,
        )

        result = reconcile_branch_financial_movements(self.branch, self.manager, repair=True)

        donation.refresh_from_db()
        transfer.refresh_from_db()
        self.assertGreaterEqual(result["created"], 2)
        self.assertIsNotNone(donation.cash_movement_id)
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch,
                source=BranchCashMovement.SOURCE_DONATION,
                source_reference=f"donation:{donation.pk}",
                amount=75000,
            ).exists()
        )
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch,
                source=BranchCashMovement.SOURCE_BANK_TRANSFER,
                source_reference=f"bank_transfer:{transfer.pk}",
                amount=50000,
                movement_type=BranchCashMovement.TYPE_OUT,
            ).exists()
        )


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerFinancialReportViewTests(TestCase):
    """Rendu de la page rapport financier et export PDF."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_report", groups=["gestionnaire"], branch=self.branch)
        self.period_month = date.today().replace(day=1)
        _login(self.client, self.manager)
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_STUDENT_PAYMENT,
            amount=180000,
            label="Paiement etudiant - Test",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_IN,
            source=BranchCashMovement.SOURCE_DONATION,
            amount=25000,
            label="Don de test",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_OUT,
            source=BranchCashMovement.SOURCE_PAYROLL,
            amount=50000,
            label="Salaire test",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_OUT,
            source=BranchCashMovement.SOURCE_BANK_TRANSFER,
            amount=30000,
            label="Versement bancaire test",
            movement_date=self.period_month,
            created_by=self.manager,
        )

    def test_report_section_renders_financial_summary(self):
        url = reverse("accounts:manager_dashboard")
        response = self.client.get(
            url,
            {
                "section": "rapport",
                "report_month": str(self.period_month.month),
                "report_year": str(self.period_month.year),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Recettes totales")
        self.assertContains(response, "Versement bancaire recommande")
        self.assertContains(response, "Paiements etudiants")

    def test_report_pdf_export_returns_pdf(self):
        url = reverse("accounts:manager_export_report_pdf")
        response = self.client.get(
            url,
            {
                "section": "rapport",
                "report_month": str(self.period_month.month),
                "report_year": str(self.period_month.year),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_report_period_controls_preserve_non_month_preset(self):
        response = self.client.get(
            reverse("accounts:manager_dashboard"),
            {"section": "rapport", "report_period": "week"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["report_period"]["preset"], "week")
        self.assertContains(response, 'id="manager-report-month" aria-label="Mois du rapport" disabled')
        html = response.content.decode()
        pdf_link = re.search(r'href="([^"]*manager/export/report/pdf/[^"]*)"', html)
        self.assertIsNotNone(pdf_link)
        self.assertNotIn("report_month=", unescape(pdf_link.group(1)))

    def test_custom_period_exposes_date_inputs(self):
        response = self.client.get(
            reverse("accounts:manager_dashboard"),
            {
                "section": "rapport",
                "report_period": "custom",
                "report_start": "2026-06-01",
                "report_end": "2026-06-15",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="manager-report-start" value="2026-06-01"')
        self.assertContains(response, 'id="manager-report-end" value="2026-06-15"')

    def test_caisse_identifies_bank_transfers_as_outflows(self):
        response = self.client.get(reverse("accounts:manager_dashboard"), {"section": "caisse"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Versements bancaires comptabilisés comme sorties de caisse")
        self.assertEqual(response.context["report_bank_transfer_total"], 30000)

    def test_closure_form_is_visibly_blocked_on_financial_anomaly(self):
        Donation.objects.create(
            branch=self.branch,
            donor_name="Don sans mouvement",
            amount=10000,
            date=self.period_month,
        )
        response = self.client.get(
            reverse("accounts:manager_dashboard"),
            {
                "section": "cloture",
                "report_month": self.period_month.month,
                "report_year": self.period_month.year,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Clôture bloquée : anomalies financières à corriger")
        self.assertContains(response, 'disabled aria-disabled="true"')

    def test_negative_cash_is_reported_as_danger(self):
        BranchCashMovement.objects.create(
            branch=self.branch,
            movement_type=BranchCashMovement.TYPE_OUT,
            source=BranchCashMovement.SOURCE_ADJUSTMENT,
            amount=500000,
            label="Anomalie caisse negative",
            movement_date=self.period_month,
            created_by=self.manager,
        )
        response = self.client.get(
            reverse("accounts:manager_dashboard"),
            {
                "section": "rapport",
                "report_month": self.period_month.month,
                "report_year": self.period_month.year,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertLess(response.context["report_available_cash_balance"], 0)
        self.assertTrue(
            any(alert["title"] == "Caisse negative" for alert in response.context["report_alerts"])
        )


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerCashSessionWorkflowTests(TestCase):
    """Workflow sessions caisse : creer, finaliser, annuler."""

    def setUp(self):
        self.branch = _create_branch()
        self.manager = _create_user("mgr_cash", groups=["gestionnaire"], branch=self.branch)
        self.programme = _create_programme()
        self.candidature = _create_candidature(self.programme, self.branch, status="accepted")
        self.inscription = Inscription.objects.create(
            candidature=self.candidature,
            amount_due=500000,
            amount_paid=0,
        )
        self.inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        self.inscription.save(update_fields=["status"])
        self.agent = PaymentAgent.objects.create(
            user=self.manager, branch=self.branch, agent_code="AGT002", is_active=True,
        )
        _login(self.client, self.manager)

    def test_cash_session_create(self):
        url = reverse("accounts:htmx_manager_cash_session_create", args=[self.inscription.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(
            CashPaymentSession.objects.filter(inscription=self.inscription, is_used=False).exists()
        )

    def test_cash_session_complete_creates_payment_and_cash_movement(self):
        session = CashPaymentSession.objects.create(
            inscription=self.inscription,
            agent=self.agent,
            verification_code="123456",
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        url = reverse("accounts:htmx_manager_cash_session_complete", args=[session.id])
        response = self.client.post(url, {"amount": "100000"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        session.refresh_from_db()
        self.assertTrue(session.is_used)
        self.assertTrue(
            Payment.objects.filter(
                inscription=self.inscription, status=Payment.STATUS_VALIDATED, amount=100000,
            ).exists()
        )
        self.assertTrue(
            BranchCashMovement.objects.filter(
                branch=self.branch, source=BranchCashMovement.SOURCE_STUDENT_PAYMENT, amount=100000,
            ).exists()
        )

    def test_cash_session_cancel(self):
        session = CashPaymentSession.objects.create(
            inscription=self.inscription,
            agent=self.agent,
            verification_code="654321",
            expires_at=timezone.now() + timedelta(minutes=15),
        )
        url = reverse("accounts:htmx_manager_cash_session_cancel", args=[session.id])
        response = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        session.refresh_from_db()
        self.assertTrue(session.is_used)


def _create_positioning_fee(programme, level="L1", amount=100000):
    from formations.models import Fee, ProgrammeYear

    year_number = int(level[1:]) if len(level) > 1 and level[1:].isdigit() else 1
    programme_year = ProgrammeYear.objects.create(programme=programme, year_number=year_number)
    Fee.objects.create(programme_year=programme_year, label="Frais inscription", amount=amount, due_month="octobre")
    return programme_year


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerCouponWorkflowTests(TestCase):
    """Workflow coupons cote gestionnaire : application a la creation et sur inscription existante."""

    def setUp(self):
        from coupons.models import Coupon

        self.Coupon = Coupon
        self.branch = _create_branch()
        self.manager = _create_user("mgr_coupon", groups=["gestionnaire"], branch=self.branch)
        self.programme = _create_programme()
        _create_positioning_fee(self.programme, level="L1", amount=100000)
        _login(self.client, self.manager)

    def _create_coupon(self, **kw):
        defaults = dict(
            code="MASTER-DG-2026",
            label="Coupon test",
            discount_type=self.Coupon.DISCOUNT_PERCENTAGE,
            value=20,
            valid_from=timezone.now() - timedelta(days=1),
            valid_until=timezone.now() + timedelta(days=30),
            max_redemptions=1,
        )
        defaults.update(kw)
        return self.Coupon.objects.create(**defaults)

    def test_inscription_create_with_valid_coupon_reduces_amount(self):
        coupon = self._create_coupon()
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        academic_class = _create_academic_class(self.programme, self.branch, level="L1")
        url = reverse("accounts:htmx_inscription_create", args=[candidature.id])
        response = self.client.post(
            url,
            {"academic_level": "L1", "academic_class": str(academic_class.pk), "coupon_code": coupon.code},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        inscription = Inscription.objects.get(candidature=candidature)
        self.assertEqual(inscription.amount_due, 80000)
        self.assertIn("Montant apres reduction", response.headers["HX-Trigger"])

    def test_inscription_create_with_invalid_coupon_still_creates_inscription(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        academic_class = _create_academic_class(self.programme, self.branch, level="L1")
        url = reverse("accounts:htmx_inscription_create", args=[candidature.id])
        response = self.client.post(
            url,
            {"academic_level": "L1", "academic_class": str(academic_class.pk), "coupon_code": "DOESNOTEXIST"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        inscription = Inscription.objects.get(candidature=candidature)
        self.assertEqual(inscription.amount_due, 100000)
        self.assertIn('"type": "warning"', response.headers["HX-Trigger"])

    def test_apply_coupon_on_existing_inscription(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000, amount_paid=0)
        inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        inscription.save(update_fields=["status"])
        coupon = self._create_coupon(discount_type=self.Coupon.DISCOUNT_FIXED, value=25000)

        url = reverse("accounts:htmx_inscription_apply_coupon", args=[inscription.id])
        response = self.client.post(url, {"coupon_code": coupon.code}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        inscription.refresh_from_db()
        self.assertEqual(inscription.amount_due, 75000)
        trigger = json.loads(response.headers["HX-Trigger"])
        self.assertEqual(trigger["couponApplied"]["inscription_id"], inscription.id)
        self.assertEqual(trigger["couponApplied"]["amount_after"], 75000)
        self.assertContains(response, f'id="inscription-balance-{inscription.id}"')

    def test_apply_coupon_rejected_after_full_payment(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000, amount_paid=0)
        inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        inscription.save(update_fields=["status"])
        Payment.objects.create(inscription=inscription, amount=100000, method=Payment.METHOD_CASH, status=Payment.STATUS_VALIDATED)
        inscription.update_financial_state()
        coupon = self._create_coupon()

        url = reverse("accounts:htmx_inscription_apply_coupon", args=[inscription.id])
        response = self.client.post(url, {"coupon_code": coupon.code}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        inscription.refresh_from_db()
        self.assertEqual(inscription.amount_due, 100000)
        self.assertContains(response, "entièrement payée")

    def test_coupon_preview_valid_code(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000, amount_paid=0)
        inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        inscription.save(update_fields=["status"])
        coupon = self._create_coupon()

        url = reverse("accounts:htmx_coupon_preview")
        response = self.client.get(
            url, {"inscription_id": inscription.id, "code": coupon.code}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-coupon-valid="1"')

    def test_coupon_preview_invalid_code(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000, amount_paid=0)
        inscription.status = Inscription.STATUS_AWAITING_PAYMENT
        inscription.save(update_fields=["status"])

        url = reverse("accounts:htmx_coupon_preview")
        response = self.client.get(
            url, {"inscription_id": inscription.id, "code": "DOESNOTEXIST"}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-coupon-valid="0"')

    def test_expired_coupon_message_is_clear(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000)
        coupon = self._create_coupon(
            valid_from=timezone.now() - timedelta(days=10),
            valid_until=timezone.now() - timedelta(days=1),
        )
        response = self.client.get(
            reverse("accounts:htmx_coupon_preview"),
            {"inscription_id": inscription.id, "code": coupon.code},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ce coupon a expiré")

    def test_coupon_apply_is_post_only(self):
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000)
        response = self.client.get(
            reverse("accounts:htmx_inscription_apply_coupon", args=[inscription.id])
        )
        self.assertEqual(response.status_code, 405)

    def test_coupon_preview_rejects_cross_branch_inscription(self):
        other_branch = _create_branch("CPN2", "Autre annexe coupon")
        candidature = _create_candidature(self.programme, other_branch, status="accepted")
        inscription = Inscription.objects.create(candidature=candidature, amount_due=100000)
        coupon = self._create_coupon()

        response = self.client.get(
            reverse("accounts:htmx_coupon_preview"),
            {"inscription_id": inscription.id, "code": coupon.code},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 404)

    def test_coupon_preview_rejects_invalid_inscription_identifier(self):
        response = self.client.get(
            reverse("accounts:htmx_coupon_preview"),
            {"inscription_id": "invalid", "code": "ANY"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Inscription invalide", status_code=400)

    def test_active_coupon_widget_is_scoped_to_manager_branch(self):
        visible = self._create_coupon(code="VISIBLE-MANAGER", max_redemptions=5)
        visible.branches.add(self.branch)
        other_branch = _create_branch("CPN3", "Annexe coupon cache")
        hidden = self._create_coupon(code="HIDDEN-MANAGER", max_redemptions=5)
        hidden.branches.add(other_branch)

        response = self.client.get(reverse("accounts:widget_active_coupons"), HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, visible.code)
        self.assertNotContains(response, hidden.code)

        dashboard_response = self.client.get(
            reverse("accounts:manager_dashboard"), {"section": "overview"}
        )
        self.assertEqual(dashboard_response.status_code, 200)
        self.assertContains(dashboard_response, visible.code)


@override_settings(
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ManagerContextualDocumentTests(TestCase):
    """Phase 3 : documents proteges par role et annexe."""

    def setUp(self):
        self.branch = _create_branch("DOC", "Annexe Documents")
        self.other_branch = _create_branch("OTH", "Autre Annexe")
        self.manager = _create_user("mgr_docs", groups=["gestionnaire"], branch=self.branch)
        self.programme = _create_programme()
        candidature = _create_candidature(self.programme, self.branch, status="accepted")
        self.inscription = Inscription.objects.create(candidature=candidature, amount_due=100000)
        other_candidature = _create_candidature(
            self.programme,
            self.other_branch,
            status="accepted",
            email="autre@test.com",
        )
        self.other_inscription = Inscription.objects.create(
            candidature=other_candidature,
            amount_due=100000,
        )
        self.employee = _create_user("employee_docs", branch=self.branch, position="secretary")
        self.payroll = PayrollEntry.objects.create(
            branch=self.branch,
            employee=self.employee,
            period_month=date(2026, 7, 1),
            base_salary=150000,
        )
        self.donation = Donation.objects.create(
            branch=self.branch,
            donor_name="Partenaire Test",
            amount=50000,
            receipt_number="DON-DOC-001",
        )
        self.expense = BranchExpense.objects.create(
            branch=self.branch,
            title="Depense sans justificatif",
            amount=25000,
        )
        _login(self.client, self.manager)

    @patch("accounts.dashboards.htmx_global.generate_esfe_pdf", return_value=b"%PDF fiche")
    def test_inscription_sheet_is_generated_for_manager_branch(self, _generate_pdf):
        response = self.client.get(
            reverse("accounts:manager_inscription_sheet_pdf", args=[self.inscription.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("fiche-inscription-", response["Content-Disposition"])

    def test_cross_branch_document_is_not_visible(self):
        response = self.client.get(
            reverse("accounts:manager_inscription_sheet_pdf", args=[self.other_inscription.pk])
        )
        self.assertEqual(response.status_code, 404)

    def test_non_manager_cannot_download_document(self):
        outsider = _create_user("docs_outsider", branch=self.branch, position="secretary")
        _login(self.client, outsider)
        response = self.client.get(
            reverse("accounts:manager_inscription_sheet_pdf", args=[self.inscription.pk])
        )
        self.assertEqual(response.status_code, 403)

    @patch("accounts.dashboards.htmx_global.build_payroll_pdf", return_value=b"%PDF paie")
    def test_payroll_sheet_uses_existing_pdf_service(self, _build_pdf):
        response = self.client.get(
            reverse("accounts:manager_payroll_sheet_pdf", args=[self.payroll.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("fiche-paie-2026-07", response["Content-Disposition"])

    @patch("accounts.dashboards.htmx_global.build_donation_receipt", return_value=b"%PDF don")
    def test_donation_receipt_uses_existing_pdf_service(self, _build_pdf):
        response = self.client.get(
            reverse("accounts:manager_donation_receipt_pdf", args=[self.donation.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("recu-don-don-doc-001.pdf", response["Content-Disposition"])

    def test_missing_expense_pdf_is_not_downloadable(self):
        response = self.client.get(
            reverse("accounts:manager_expense_supporting_document_pdf", args=[self.expense.pk])
        )
        self.assertEqual(response.status_code, 404)
